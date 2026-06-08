"""
Stage 5: write the two-tab Excel tracker.

  Tab "M&A"            — Strategic M&A + PE buyouts
  Tab "Capital Raises" — growth raises above the disclosure floor

Headers match the precedent sheet EXACTLY. Undisclosed fields are left BLANK
(never "-", never a guess). Money columns hold real numbers in $M so they sort
and filter. Header row gets the blue house style + an autofilter dropdown.
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config
from models import DealRecord, MA_TYPES, CAPITAL_RAISE, DEAL_TYPE_LABEL, fmt_date, millions_value

MA_HEADERS = ["Sector", "Target Country", "Deal Type", "Date", "Target", "Acquirer",
              "EV ($M)", "EV / Revenue", "EV / EBITDA", "Target Description",
              "Link / Press Release"]

RAISE_HEADERS = ["Sector", "Target Country", "Date", "Target", "Lead Investor(s)",
                 "Amount ($M)", "Valuation ($M)", "EV / Revenue", "EV / EBITDA",
                 "Target Description", "Link / Press Release"]

_HEADER_FILL = PatternFill("solid", fgColor="1F6FB2")     # house blue
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_CELL_ALIGN = Alignment(vertical="center", wrap_text=True)
_THIN = Side(style="thin", color="D0D7DE")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_WIDTHS = {"Sector": 22, "Target Country": 16, "Deal Type": 14, "Date": 12,
           "Target": 22, "Acquirer": 22, "Lead Investor(s)": 26, "EV ($M)": 11,
           "Amount ($M)": 12, "Valuation ($M)": 14, "EV / Revenue": 12,
           "EV / EBITDA": 12, "Target Description": 60, "Link / Press Release": 44}


def _seg_label(seg: str | None) -> str:
    return config.SEGMENTS[seg]["label"] if seg in config.SEGMENTS else ""


def _ma_row(r: DealRecord) -> list:
    return [_seg_label(r.segment), r.target_country or "",
            DEAL_TYPE_LABEL.get(r.deal_type, r.deal_type), fmt_date(r.date_announced),
            r.target or "", r.counterparty or "", millions_value(r.ev_text),
            r.ev_revenue or "", r.ev_ebitda or "", r.target_description or "",
            r.source_url or ""]


def _raise_row(r: DealRecord) -> list:
    return [_seg_label(r.segment), r.target_country or "", fmt_date(r.date_announced),
            r.target or "", r.counterparty or "", millions_value(r.amount_text),
            millions_value(r.valuation_text), r.ev_revenue or "", r.ev_ebitda or "",
            r.target_description or "", r.source_url or ""]


def _write_sheet(ws, headers: list[str], rows: list[list]):
    ws.append(headers)
    for c, name in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font, cell.alignment, cell.border = (
            _HEADER_FILL, _HEADER_FONT, _HEADER_ALIGN, _BORDER)
        ws.column_dimensions[get_column_letter(c)].width = _WIDTHS.get(name, 16)
    money_cols = {i + 1 for i, h in enumerate(headers)
                  if h in ("EV ($M)", "Amount ($M)", "Valuation ($M)")}
    for row in rows:
        ws.append(row)
        rr = ws.max_row
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=rr, column=c)
            cell.alignment, cell.border = _CELL_ALIGN, _BORDER
            if c in money_cols and isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0'
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"
    ws.row_dimensions[1].height = 30


def export(selected: list[DealRecord], path: str) -> tuple[str, int, int]:
    ma = [r for r in selected if r.deal_type in MA_TYPES]
    raises = [r for r in selected if r.deal_type == CAPITAL_RAISE]
    ma.sort(key=lambda r: (_seg_label(r.segment), r.date_announced or ""))
    raises.sort(key=lambda r: (_seg_label(r.segment), r.date_announced or ""))

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "M&A"
    _write_sheet(ws1, MA_HEADERS, [_ma_row(r) for r in ma])
    ws2 = wb.create_sheet("Capital Raises")
    _write_sheet(ws2, RAISE_HEADERS, [_raise_row(r) for r in raises])
    wb.save(path)
    return path, len(ma), len(raises)
